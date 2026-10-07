"""
report_xlsx — Excel ixracı (openpyxl): hesablama nəticəsi (run), müqayisə və/və ya seçilmiş çıxış CSV-ləri ayrı
vərəqlərə; «Məlumat» vərəqində ssenari, vintaj id-ləri, izah mətni, yaradılma vaxtı. Başlıqlar Azərbaycan dilində:
panel/i18n/columns_az.csv (varsa) + daxili lüğət; hər başlıq xanasının şərhində sütun kodu.
(RiskUnit/api/report_xlsx.py-dən uyğunlaşdırılıb.)
"""
import csv, io, json, re
from pathlib import Path

from apicore import ApiError, now_iso

LABELS = {"scenario": "Ssenari", "scenario_name": "Ssenarinin adı", "engine": "Mühərrik", "horizon": "Üfüq",
          "indicator": "Göstərici", "label_az": "Ad", "unit": "Vahid", "year": "İl", "baseline": "Baza", "value": "Ssenari",
          "delta": "Fərq", "delta_pct": "Fərq, % (f.b.)", "method": "Metod", "tier": "Sübut səviyyəsi", "group": "Qrup",
          "note_az": "Qeyd", "effect": "Təsir", "effect_unit": "Təsirin vahidi", "years": "İllər", "source_engine": "Mənbə",
          "qısa": "Qısa müddət", "orta": "Orta müddət", "uzun": "Uzun müddət", "spread": "Metodlararası fərq",
          "sign_agree": "İşarə uyğunluğu", "explanation_az": "İzah", "rank": "Yer", "score": "Bal",
          "kpis_used": "İstifadə olunan KPI", "kpis_missing": "Hesablanmayan KPI", "cost_mln_azn": "Xərc, mln AZN",
          "score_per_bn_azn": "Bal / 1 mlrd AZN", "kpi": "KPI", "name_az": "Ad", "direction": "İstiqamət", "weight": "Çəki",
          "norm": "Normallaşdırılmış bal", "severity_az": "Ciddilik", "family": "Ailə",
          "affected_group": "Təsirlənən qrup", "affected_sector": "Təsirlənən sektor", "proposal_az": "Təklif",
          "responsible_az": "Məsul qurum", "variant_name_az": "Yan təsir ssenarisi", "sector": "Sektor"}
MAX_ROWS = 100000


def column_labels(root):
    out = dict(LABELS)
    p = Path(root) / "panel" / "i18n" / "columns_az.csv"
    try:
        with open(p, encoding="utf-8", newline="") as f:
            out.update({r["key"]: r["az"] for r in csv.DictReader(f) if r.get("key") and r.get("az")})
    except (OSError, KeyError):
        pass
    return out


def _sheet_name(s, used):
    s = re.sub(r"[\[\]\*\?/\\:]", "_", str(s))[:31] or "vərəq"
    base, i = s, 2
    while s in used:
        s = base[:28] + "_%d" % i
        i += 1
    used.add(s)
    return s


def _cell(v):
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)[:32000]
    if isinstance(v, float) and v != v:
        return None
    return v


def run_tables(res, prefix=""):
    se = res.get("side_effects") or {}
    t = [("Əsas göstəricilər", res.get("headline_wide")), ("Üfüq üzrə təsir", res.get("headline")),
         ("Metod müqayisəsi", res.get("comparison")), ("Sektorlar", res.get("sectors")), ("Sosial", res.get("social")),
         ("Yan təsirlər", se.get("items")), ("Yumşaltma", se.get("mitigation")),
         ("Risk profili", (se.get("risk_profile") or {}).get("rows")), ("KPI", (res.get("kpi") or {}).get("values")),
         ("Mühərriklər", res.get("engines")), ("Bütün təsirlər", res.get("effects"))]
    return [(prefix + n, v) for n, v in t if v]


def compare_tables(cmp):
    return [(n, cmp.get(k)) for n, k in (("Reytinq", "ranking"), ("KPI matrisi", "matrix"), ("KPI dəyərləri", "values"),
                                         ("Əsas göstəricilər (müq.)", "headline"), ("Yan təsirlər (müq.)", "side_effects"))
            if cmp.get(k)]


def build(views, root, files=None, run=None, compare=None, title=None, meta=None):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    info = wb.active
    info.title = "Məlumat"
    used = {"Məlumat"}
    rows = [("Hesabat", title or "MİİS §15.5.4 — iqtisadi siyasətlərin təsir analizi"), ("Yaradıldı (UTC)", now_iso())]
    rows += [(k, v if isinstance(v, str) else json.dumps(v, ensure_ascii=False)) for k, v in (meta or {}).items()]
    tables = []
    if run:
        s = run.get("scenario") or {}
        rows += [("Ssenari", "%s — %s" % (s.get("id"), s.get("name_az"))), ("Başlanğıc il", s.get("start_year")),
                 ("Vintaj", (run.get("vintage") or {}).get("vintage_id")), ("İzah", run.get("text_az"))]
        tables += run_tables(run)
    if compare:
        rows += [("Müqayisə", ", ".join(x["scenario"] for x in compare.get("scenarios", []))), ("Reytinq izahı", compare.get("text_az"))]
        tables += compare_tables(compare)
    for f in files or []:
        rel, p = views.output_path(f)
        if p.suffix.lower() != ".csv":
            raise ApiError(400, "bad_parameter", "Excel-ə yalnız CSV faylları ixrac olunur: %s" % rel)
        df = views.frame(rel).head(MAX_ROWS)
        tables.append((rel.rsplit(".", 1)[0], df.astype(object).where(df.notna(), None).to_dict("records")))
        ent = views.catalog_entry(rel)
        if ent:
            rows.append((rel, ent.get("description_az", "")))
    if not tables:
        raise ApiError(400, "empty_report", "İxrac üçün cədvəl yoxdur (run_id, compare və ya files verin)")
    for k, v in rows:
        info.append([k, v])
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 26, 120
    for c in info["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDE6F0")
    labels = column_labels(root)
    for name, recs in tables:
        recs = [r for r in recs if isinstance(r, dict)][:MAX_ROWS]
        if not recs:
            continue
        cols = []
        for r in recs:
            cols += [c for c in r if c not in cols]
        ws = wb.create_sheet(_sheet_name(name, used))
        hdr, seen = [], set()
        for c in cols:
            h = labels.get(str(c), str(c))
            h = "%s (%s)" % (h, c) if h in seen else h
            seen.add(h)
            hdr.append(h)
        ws.append(hdr)
        for c, code in zip(ws[1], cols):
            c.font, c.fill = bold, fill
            if str(code) != c.value:
                c.comment = Comment("kod: %s" % code, "MİİS")
        for r in recs:
            ws.append([_cell(r.get(c)) for c in cols])
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(hdr[:60], start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(len(str(h)) + 2, 10), 45)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
