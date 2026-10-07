"""Ministry CAEM `AZE Model` reduced form  y_t = Bc + B1 y_{t-1} + B2 e_t  (48 states, annual).

COPIED (not imported) from RiskUnit/riskunit/caem_model.py (load/simulate/md5/vintage_status, 2026-10-06)
so that the PolicyUnit never imports RiskUnit in-process. Reads the pinned copy
PolicyUnit/data/ministry/CAEM.xlsx (md5 pinned, read-only, openpyxl data_only). Deviations from
baseline only; labelled "Nazirlik CAEM modeli — müqayisə" (calibrated parameters -> tier D).
"""
from __future__ import annotations

import hashlib
from functools import lru_cache

import numpy as np
import pandas as pd

from . import config

LABEL_AZ = "Nazirlik CAEM modeli — müqayisə"
N_STATE = 48
H = 12                                            # horizons 0..12 (h=1 <-> 2026)
RANGES = {"B1": (211, 258), "B2": (262, 309), "Ac": (159, 206)}
DEBT = ("d_y", "dd_y", "df_y")


def md5(path) -> str:
    h = hashlib.md5()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def workbook_path():
    for p in (config.CAEM_COPY, config.CAEM_RISK_COPY, config.CAEM_UPSTREAM):
        if p.exists():
            return p
    raise FileNotFoundError("CAEM.xlsx tapılmadı (data/ministry/CAEM.xlsx)")


def vintage_status() -> dict:
    p = workbook_path()
    out = {"copy": str(p), "copy_md5": md5(p),
           "upstream_md5": md5(config.CAEM_UPSTREAM) if config.CAEM_UPSTREAM.exists() else None}
    out["pinned_ok"] = out["copy_md5"] == config.CAEM_MD5
    out["upstream_changed"] = out["upstream_md5"] is not None and out["upstream_md5"] != out["copy_md5"]
    return out


@lru_cache(maxsize=1)
def _wb():
    import openpyxl
    return openpyxl.load_workbook(workbook_path(), read_only=True, data_only=True)


def _num(v):
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    return 0.0 if v is None or v == "" else np.nan


def _block(sheet, r0, r1, c0, c1):
    rows = _wb()[sheet].iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True)
    return np.array([[_num(v) for v in r] for r in rows], float)


def _cells(sheet, r0, r1, c0, c1):
    return [[getattr(v, "text", v) for v in r] for r in
            _wb()[sheet].iter_rows(min_row=r0, max_row=r1, min_col=c0, max_col=c1, values_only=True)]


@lru_cache(maxsize=1)
def load() -> dict:
    m = {k: _block("AZE Model", a, b, 3, 50) for k, (a, b) in RANGES.items() if k != "Ac"}
    a, b = RANGES["Ac"]
    m["Ac"] = np.nan_to_num(_block("AZE Model", a, b, 3, 3)[:, 0])
    m["states"] = pd.DataFrame(_cells("8a. Simulation", 62, 109, 1, 3), columns=["name", "unit", "code"])
    m["psi"] = float(_cells("Parametrization", 26, 26, 7, 7)[0][0])     # G26 domestic debt share
    m["ix"] = {c: i for i, c in enumerate(m["states"]["code"])}
    return m


def simulate(E=None, targets=None, fix_debt_bug: bool = False, m=None) -> np.ndarray:
    """Deviation-from-baseline paths (48 x 13) exactly as `8a. Simulation` computes them.
    E: shocks (rows = 8a rows 5..52 = state order); targets: {state: array(13)} imposed paths
    (NaN = free), solved for the own shock each period."""
    m = m or load()
    B1, B2, c, ix, psi = m["B1"], m["B2"], m["Ac"], m["ix"], m["psi"]
    E = np.zeros((N_STATE, H + 1)) if E is None else np.array(E, float).copy()
    y = np.zeros((N_STATE, H + 1))
    nd = N_STATE - len(DEBT)
    tg = {ix[k]: np.asarray(v, float) for k, v in (targets or {}).items()}
    for t in range(1, H + 1):
        act = [i for i, p in tg.items() if np.isfinite(p[t])]
        if act:
            base = c + B1 @ y[:, t - 1] + B2 @ E[:, t]
            gap = np.array([tg[i][t] for i in act]) - base[act]
            E[act, t] += np.linalg.solve(B2[np.ix_(act, act)], gap)
        y[:nd, t] = (c + B1 @ y[:, t - 1] + B2 @ E[:, t])[:nd]
        den = 1 + y[ix["dy"], t - 1] / 100 + y[ix["dP"], t - 1] / 100
        y[ix["dd_y"], t] = ((1 + y[ix["CR"], t - 1] / 100) / den * y[ix["dd_y"], t - 1]
                            - psi * y[ix["pb_y"], t])
        lag = y[ix["df_y"], t - 1] if (t == 1 or fix_debt_bug) else y[ix["dd_y"], t - 1]
        y[ix["df_y"], t] = ((1 + y[ix["R_f"], t] / 100) * (1 + y[ix["dS"], t] / 100) / den * lag
                            - (1 - psi) * y[ix["g_y"], t])
        y[ix["d_y"], t] = y[ix["dd_y"], t] + y[ix["df_y"], t]
    return y


def zero_check(m=None) -> float:
    """Max |response| with no shocks (the constant Bc is part of the deviation system)."""
    return float(np.nanmax(np.abs(simulate(m=m))))
