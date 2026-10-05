"""
report_xlsx — sadə Excel hesabatı (openpyxl): göstəricilər × illər × ssenarilər, tənliklərin xülasəsi və
(istəyə görə) saxlanmış ssenarilərin nəticələri. Əsas hesabat qurucusu brauzerdədir; bu, server tərəfi nüsxədir.

Spesifikasiya (POST /api/v1/reports):
  {"title": "...", "indicators": ["fr1:rgdp", ...] | "module": "FR4", "scenarios": ["Baseline", ...],
   "from": 2015, "to": 2030, "equations": true | ["FR1.C3_man", ...], "saved_scenarios": ["s2026..."]}
"""
import io
from datetime import datetime

from apicore import ApiError, SCEN_AZ, SCENARIOS, module_name, now_iso

HDR_FILL = "1F3B5A"


def _style_header(ws, row, ncol):
    from openpyxl.styles import Font, PatternFill, Alignment
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HDR_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _flatten_series(obj, prefix=""):
    """Mühərrik nəticəsində bütün {"series": {id: {il: dəyər}}} bloklarını tapır."""
    out = {}
    if isinstance(obj, dict):
        s = obj.get("series")
        if isinstance(s, dict):
            for k, v in s.items():
                if isinstance(v, dict):
                    out[(prefix, k)] = v
        for k, v in obj.items():
            if k != "series" and isinstance(v, dict):
                out.update(_flatten_series(v, k if not prefix else "%s/%s" % (prefix, k)))
    return out


def build(spec, views, saved_get=None, vintage=None):
    import openpyxl
    from openpyxl.utils import get_column_letter
    if not isinstance(spec, dict):
        raise ApiError(400, "bad_request", "Hesabat spesifikasiyası JSON obyekt olmalıdır")
    ids = spec.get("indicators") or []
    module = module_name(spec.get("module"), required=False)
    if not ids and not module:
        raise ApiError(400, "bad_request", "«indicators» siyahısı və ya «module» tələb olunur")
    if len(ids) > 2000:
        raise ApiError(413, "too_large", "Bir hesabatda maksimum 2000 göstərici")
    scen = spec.get("scenarios") or SCENARIOS
    bad = [s for s in scen if s not in SCENARIOS + ["Actual"]]
    if bad:
        raise ApiError(400, "bad_parameter", "Naməlum ssenari: %s" % ", ".join(bad))
    y0, y1 = spec.get("from"), spec.get("to")
    fc = views.forecasts(module=module, ids=ids or None, scenarios=scen, y0=y0, y1=y1)
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = "Məlumat"
    title = spec.get("title") or "MikroModel — proqnoz hesabatı"
    rows = [("Hesabat", title), ("Yaradılıb", now_iso()), ("Mənbə", fc["source"]),
            ("Ssenarilər", ", ".join("%s (%s)" % (SCEN_AZ.get(s, s), s) for s in scen)),
            ("Dövr", "%s – %s" % (y0 or "…", y1 or "…")), ("Göstəricilərin sayı", len(fc["series"]))]
    for k, v in (vintage or {}).items():
        rows.append((k, str(v)))
    for w in fc.get("warnings", [])[:50]:
        rows.append(("Xəbərdarlıq", w.get("message")))
    for r in rows:
        about.append(list(r))
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 100
    # ---- indicators x years x scenarios
    ws = wb.create_sheet("Göstəricilər")
    years = sorted({it["period"] for it in fc["items"]})
    head = ["ID", "Modul", "Göstərici", "Vahid", "Ssenari"] + years
    ws.append(head)
    _style_header(ws, 1, len(head))
    grid = {}
    for it in fc["items"]:
        grid.setdefault((it["id"], it["scenario"]), {})[it["period"]] = it["value"]
    order = {s: i for i, s in enumerate(["Actual"] + SCENARIOS)}
    for (sid, s) in sorted(grid, key=lambda k: (k[0], order.get(k[1], 9))):
        m = fc["series"].get(sid, {})
        ws.append([sid, m.get("module"), m.get("label_az"), m.get("unit_az"), SCEN_AZ.get(s, s)] + [grid[(sid, s)].get(y) for y in years])
    for i, w in enumerate([34, 7, 48, 14, 10] + [11] * len(years), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "F2"
    # ---- equations
    eqsel = spec.get("equations")
    if eqsel:
        mods = sorted({m.get("module") for m in fc["series"].values() if m.get("module")}) or ([module] if module else [])
        eqs = []
        if isinstance(eqsel, list):
            for e in eqsel:
                eqs += views.equations(eq_id=str(e), full=True)["items"]
        else:
            for m in mods:
                eqs += views.equations(module=m, full=True)["items"]
        we = wb.create_sheet("Tənliklər")
        h = ["Tənlik", "Modul", "Başlıq", "Asılı dəyişən", "Qiymətləndirici", "Kovariasiya", "Müşahidələr", "R²",
             "Düzəldilmiş R²", "Proqnozda istifadə", "Dayanıqlıq"]
        we.append(h)
        _style_header(we, 1, len(h))
        for e in eqs:
            fit, smp = e.get("fit") or {}, e.get("sample") or {}
            we.append([e.get("id"), e.get("module"), e.get("title_az"), (e.get("dependent") or {}).get("label_az") or
                       (e.get("dependent") or {}).get("code"), e.get("estimator"), e.get("cov_type"), smp.get("n"),
                       fit.get("r2"), fit.get("r2_adj"), "bəli" if e.get("used_in_forecast") else "xeyr",
                       (e.get("robustness") or {}).get("verdict")])
        we.append([])
        h2 = ["Tənlik", "Əmsal", "Ad", "Qiymət", "Standart xəta", "t-statistikası", "p-dəyəri", "95% EI aşağı", "95% EI yuxarı",
              "İstifadə olunan", "Sabit"]
        we.append(h2)
        _style_header(we, we.max_row, len(h2))
        for e in eqs:
            for c in e.get("coefficients") or []:
                we.append([e.get("id"), c.get("name"), c.get("label_az"), c.get("coef"), c.get("se"), c.get("t"), c.get("p"),
                           c.get("ci_low"), c.get("ci_high"), c.get("used_value"), "bəli" if c.get("fixed") else ""])
        for i, w in enumerate([22, 14, 40, 14, 14, 14, 10, 12, 12, 14, 14], start=1):
            we.column_dimensions[get_column_letter(i)].width = w
        if not eqs:
            we.append(["Tənlik reyestri hələ yaradılmayıb (output/FRx_equations.json)"])
    # ---- saved scenarios
    for sid in spec.get("saved_scenarios") or []:
        sv = saved_get(sid)
        wsx = wb.create_sheet(("Ssenari %s" % sv["name"])[:31].replace("/", "-").replace(":", "-"))
        wsx.append(["Ad", sv["name"]]); wsx.append(["Müəllif", sv.get("author")]); wsx.append(["Ssenari", sv.get("scenario")])
        wsx.append(["Yenilənib", sv.get("updated_at")]); wsx.append([])
        ser = _flatten_series(sv.get("result") or {})
        ys = sorted({int(y) for v in ser.values() for y in v if str(y).isdigit()})
        wsx.append(["Blok", "Göstərici"] + ys)
        _style_header(wsx, wsx.max_row, 2 + len(ys))
        for (blk, k), v in sorted(ser.items()):
            wsx.append([blk, k] + [v.get(str(y), v.get(y)) for y in ys])
        if not ser:
            wsx.append(["Bu ssenaridə saxlanmış nəticə yoxdur"])
    buf = io.BytesIO()
    wb.save(buf)
    name = "MikroModel_hesabat_%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M%S")
    return buf.getvalue(), name
