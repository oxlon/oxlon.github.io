"""
validate_workbook — yeni «Statistik data dinamika *.xlsx» iş kitabının yoxlanılması.

Dəftərlərin istinad etdiyi hər (vərəq, sətir, ad) ünvanı dəftərlərin öz kodundan çıxarılır (nbextract):
  * FR1 VARMAP  — (vərəq, sətir, ad fraqmenti, vahid fraqmenti); FR1-in öz yoxlaması təkrarlanır
  * FR3 WAGEMAP — (vərəq, sətir, ad fraqmenti); FR3-ün öz yoxlaması təkrarlanır
  * FR4/FR5/FR10/FR12 `wb_rows(vərəq, {sətir: ad})` — həmin ünvandakı ad indiki iş kitabındakı adla
    müqayisə olunur; `wb_rows(vərəq, {ad: regex})` (FR12) — regex vərəqdə tapılmalıdır
Əlavə olaraq dəftərlərdə adı çəkilən bütün vərəqlər və il başlıqları yoxlanılır.
"""
import re, unicodedata
from pathlib import Path

import nbextract as X

MONTHS = {'yanvar': 1, 'fevral': 2, 'mart': 3, 'aprel': 4, 'may': 5, 'iyun': 6, 'iyul': 7, 'avqust': 8,
          'sentyabr': 9, 'oktyabr': 10, 'noyabr': 11, 'dekabr': 12}


def az_lower(s):                      # FR1 / FR3
    return str(s).lower().replace('i̇', 'i')


def az_lower_nfc(s):                  # FR4 / FR5 / FR10 / FR12
    return unicodedata.normalize('NFC', str(s).lower().replace('̇', ''))


def clean(s):
    return re.sub(r'\s+', ' ', str(s).replace('\n', ' ')).strip()


def _norm(s):                         # FR1
    return re.sub(r'[^a-z0-9əğıöşüç ]', '', az_lower(unicodedata.normalize('NFKD', str(s))))


def _key(s):
    return _norm(s).replace(' ', '')


def parse_header(h):
    """FR1: sütun başlığı → (il, son ay, A|YTD) və ya None."""
    if h is None:
        return None
    if isinstance(h, (int, float)) and not isinstance(h, bool) and float(h).is_integer() and 1985 <= h <= 2035:
        return (int(h), 12, 'A')
    s = clean(h)
    if re.fullmatch(r'(19|20)\d\d(\s*-ci il)?', s):
        return (int(s[:4]), 12, 'A')
    ym = re.search(r'(19|20)\d\d', s)
    if not ym:
        return None
    year, low = int(ym.group(0)), az_lower(s)
    hits = [(low.rfind(k), v) for k, v in MONTHS.items() if k in low]
    if hits:
        return (year, max(hits)[1], 'YTD')
    q = re.search(r'\(\s*([ivx]+)\s*rüb', low)
    return (year, {'i': 3, 'ii': 6, 'iii': 9, 'iv': 12}.get(q.group(1), 12), 'YTD') if q else None


def extract_addresses(root):
    root, out = Path(root), []
    for var, t in (X.assigned(root / 'FR1.ipynb', 'VARMAP') or {}).items():
        if isinstance(t, tuple) and len(t) >= 3:
            out.append(dict(module='FR1', var=var, sheet=t[0], row=t[1], label=t[2], unit=t[3] if len(t) > 3 else '', rule='fr1'))
    for var, t in (X.assigned(root / 'FR3.ipynb', 'WAGEMAP') or {}).items():
        if isinstance(t, tuple) and len(t) >= 3:
            out.append(dict(module='FR3', var=var, sheet=t[0], row=t[1], label=t[2], unit='', rule='fr3'))
    for m in ('FR4', 'FR5', 'FR10', 'FR12'):
        for _, args in X.calls(root / ('%s.ipynb' % m), 'wb_rows'):
            if not args or len(args) < 2 or not isinstance(args[0], str) or not isinstance(args[1], dict):
                continue
            for k, v in args[1].items():
                if isinstance(k, int):
                    out.append(dict(module=m, var=str(v), sheet=args[0], row=k, label=None, unit=None, rule='row'))
                elif isinstance(v, str):
                    out.append(dict(module=m, var=str(k), sheet=args[0], row=None, label=v, unit=None, rule='pattern'))
    return out


def named_sheets(root, sheetnames):
    """Dəftərlərin kodunda sətir sabiti kimi adı çəkilən vərəqlər (istinad iş kitabının vərəqləri arasından)."""
    found = {}
    for m in ('FR1', 'FR3', 'FR4', 'FR5', 'FR10', 'FR12'):
        p = Path(root) / ('%s.ipynb' % m)
        if not p.exists():
            continue
        consts = set(re.findall(r"'([^'\n]{2,60})'", X.source(p))) | set(re.findall(r'"([^"\n]{2,60})"', X.source(p)))
        for s in sheetnames:
            if s in consts:
                found.setdefault(s, set()).add(m)
    return found


def read_book(path, sheets):
    """{vərəq: sətirlər}, vərəq adları. Yalnız lazım olan vərəqlər oxunur (read_only)."""
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        mats = {s: [list(r) for r in wb[s].iter_rows(values_only=True)] for s in sheets if s in names}
    finally:
        wb.close()
    return names, mats


def _cell(mat, r, c):
    row = mat[r] if 0 <= r < len(mat) else []
    return row[c] if c < len(row) else None


def _first_label(mat, r):
    row = mat[r] if 0 <= r < len(mat) else []
    return next((str(v) for v in row if v is not None and not (isinstance(v, float) and v != v) and str(v).strip()), '')


def years(mat, scan=4):
    best = []
    for i in range(min(scan, len(mat))):
        ys = sorted({p[0] for p in (parse_header(v) for v in mat[i]) if p and p[2] == 'A'})
        if len(ys) > len(best):
            best = ys
    return best


def _err(lst, a, msg, found=None):
    lst.append({"module": a.get("module"), "var": a.get("var"), "sheet": a.get("sheet"), "row": a.get("row"),
                "expected": a.get("label"), "found": found, "message": msg})


def check_address(a, mat, ref_mat, errors):
    r = (a["row"] or 0) - 1
    if a["rule"] == "pattern":
        pat = re.compile(a["label"])
        if not any(row and row[0] and pat.search(az_lower_nfc(str(row[0])).strip()) for row in mat):
            _err(errors, a, "Vərəqdə «%s» şablonuna uyğun sətir tapılmadı" % a["label"])
        return
    if r >= len(mat):
        _err(errors, a, "Sətir %d vərəqdən kənardadır (vərəqdə %d sətir var)" % (a["row"], len(mat)))
        return
    row = mat[r]
    if a["rule"] == "fr1":
        names = [clean(row[c]) for c in (0, 1, 2) if c < len(row) and row[c] is not None and clean(row[c])]
        units = [clean(row[c]) for c in (1, 2, 3) if c < len(row) and row[c] is not None]
        ok_l = any(_key(a["label"]) in _key(c) for c in names)
        ok_u = (a["unit"] or '') == '' or any(_norm(a["unit"]) in _norm(u) for u in units)
        if not ok_l:
            _err(errors, a, "Ad uyğun gəlmir (FR1 VARMAP)", names[0] if names else '')
        elif not ok_u:
            _err(errors, a, "Ölçü vahidi uyğun gəlmir: «%s» gözlənilirdi" % a["unit"], units[0] if units else '')
    elif a["rule"] == "fr3":
        lab = next((clean(row[c]) for c in (0, 1, 2) if c < len(row) and row[c] is not None and clean(row[c])), '')
        if az_lower(a["label"]).replace(' ', '') not in az_lower(lab).replace(' ', ''):
            _err(errors, a, "Ad uyğun gəlmir (FR3 WAGEMAP)", lab[:80])
    else:                                                       # 'row': compare with the installed workbook
        lab = _first_label(mat, r)
        if not lab:
            _err(errors, a, "Ünvan boşdur — dəftər bu sətri oxuyur")
        elif ref_mat is not None:
            ref = _first_label(ref_mat, r)
            if ref and _key(ref) != _key(lab):
                a = dict(a, label=ref[:80])
                _err(errors, a, "Bu ünvanda indiki iş kitabındakından fərqli göstərici var", lab[:80])


def validate_workbook(path, root, reference=None, max_errors=300):
    """Hesabat: {"ok", "errors", "warnings", "summary"}. reference — indi quraşdırılmış iş kitabı (varsa)."""
    errors, warnings = [], []
    root = Path(root)
    addrs = extract_addresses(root)
    if not addrs:
        warnings.append({"message": "Dəftərlərdən ünvan cədvəli çıxarıla bilmədi — yalnız vərəq adları yoxlanılır"})
    ref_names, ref_mats = [], {}
    if reference and Path(reference).exists() and Path(reference).resolve() != Path(path).resolve():
        try:
            ref_names, _ = read_book(reference, [])
        except Exception as e:                                   # pragma: no cover
            warnings.append({"message": "İndiki iş kitabı oxunmadı: %s" % e})
    named = named_sheets(root, ref_names) if ref_names else {}
    need = sorted({a["sheet"] for a in addrs} | set(named))
    try:
        names, mats = read_book(path, need)
    except Exception as e:
        return {"ok": False, "kind": "workbook", "errors": [{"message": "Fayl Excel (.xlsx) kimi açılmadı: %s" % e}],
                "warnings": warnings, "summary": {}}
    if ref_names:
        try:
            _, ref_mats = read_book(reference, sorted({a["sheet"] for a in addrs if a["rule"] == "row"}))
        except Exception:
            ref_mats = {}
    missing = sorted({a["sheet"] for a in addrs} - set(names))
    for s in missing:
        n = sum(1 for a in addrs if a["sheet"] == s)
        errors.append({"sheet": s, "message": "Vərəq yoxdur: «%s» (%d ünvan, modullar: %s)"
                       % (s, n, ", ".join(sorted({a["module"] for a in addrs if a["sheet"] == s})))})
    for s, mods in sorted(named.items()):
        if s not in names and s not in missing:
            warnings.append({"sheet": s, "message": "Dəftərlərdə adı çəkilən vərəq yoxdur: «%s» (%s)" % (s, ", ".join(sorted(mods)))})
    by_mod = {}
    for a in addrs:
        if a["sheet"] not in mats:
            continue
        before = len(errors)
        check_address(a, mats[a["sheet"]], ref_mats.get(a["sheet"]), errors)
        st = by_mod.setdefault(a["module"], {"checked": 0, "failed": 0})
        st["checked"] += 1
        st["failed"] += len(errors) - before
    yrs = {}
    for s, mat in mats.items():
        ys = years(mat)
        if ys:
            yrs[s] = [ys[0], ys[-1]]
        elif s in {a["sheet"] for a in addrs if a["rule"] in ("fr1", "fr3")}:
            errors.append({"sheet": s, "message": "Vərəqin ilk 4 sətrində il başlıqları tapılmadı"})
    last = max((v[1] for v in yrs.values()), default=None)
    if ref_names and yrs:
        try:
            _, rm = read_book(reference, list(yrs))
            for s, (_, y1) in yrs.items():
                ry = years(rm.get(s, []))
                if ry and y1 < ry[-1]:
                    warnings.append({"sheet": s, "message": "Son illik il %d indiki iş kitabındakından (%d) azdır" % (y1, ry[-1])})
        except Exception:
            pass
    fn = Path(path).name
    if not re.match(r"^Statistik data dinamika.*\.xlsx$", fn, re.I):
        warnings.append({"message": "Fayl adı «Statistik data dinamika *.xlsx» şablonuna uyğun deyil: %s" % fn})
    total = len(errors)
    return {"ok": total == 0, "kind": "workbook", "errors": errors[:max_errors], "warnings": warnings,
            "summary": {"sheets": len(names), "addresses_checked": sum(v["checked"] for v in by_mod.values()),
                        "errors_total": total, "by_module": by_mod, "years": yrs, "last_annual_year": last,
                        "reference": Path(reference).name if ref_names else None}}
