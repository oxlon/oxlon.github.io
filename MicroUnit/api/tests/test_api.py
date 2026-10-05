"""
test_api — server.py-nin əsas endpoint-ləri: sağlamlıq, nişanlar, CORS, statik fayllar və yol keçidi, yükləmələr
(xam və multipart), tətbiq (YALNIZ müvəqqəti data qovluğuna), məlumat oxuma, hesabat.

    python3 -m unittest discover -s api/tests -v
"""
import hashlib, io, json, os, shutil, sys, unittest
from pathlib import Path

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import helpers as H                                                   # noqa: E402
import multipart_form                                                 # noqa: E402


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


class ApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.real_fp = H.REAL_ROOT / "data/firm_panel"
        cls.real_state = sorted((p.name, p.stat().st_mtime_ns) for p in cls.real_fp.iterdir())
        cls.root = H.make_root()
        out = cls.root / "output"
        (out / "FR1_forecast_full.csv").write_text(
            ",rgdp,cpi,scenario\n2025,100,3.0,ACTUAL\n2026,102,3.1,Baseline\n2027,104,3.2,Baseline\n"
            "2026,101,3.5,Adverse\n2027,101.5,3.6,Adverse\n", encoding="utf-8")
        (out / "FR1_indicator_catalog.csv").write_text(
            "id,module,group_az,label_az,label_en,unit_az,freq,kind,has_forecast,scenarios,has_band,source_csv,source_column,equation_ids,imputed_years\n"
            "fr1:rgdp,FR1,Makro,Real ÜDM,Real GDP,mln manat,A,level,True,Baseline;Adverse,False,FR1_forecast_full.csv,rgdp,FR1.C1,\n"
            "fr1:cpi,FR1,Qiymət,İnflyasiya,CPI,%,A,rate,True,Baseline;Adverse,False,FR1_forecast_full.csv,cpi,,\n", encoding="utf-8")
        (cls.root / "panel/data/fr4.js").write_text(
            "/* test */\nwindow.MICRO=window.MICRO||{S:[]};window.MICRO.S=window.MICRO.S.concat("
            + json.dumps([{"i": "fr4:emp:agr", "f": "FR4", "g": "Məşğulluq", "e": "Kənd təsərrüfatı", "v": "", "u": "min nəfər",
                           "k": "lvl", "s": {"B": [1, 2, 3, 4, 5], "A": [1, 1, 1, 1, 1]}, "h": [[2024, 0.8], [2025, 0.9]]}],
                          ensure_ascii=False) + ");\n", encoding="utf-8")
        eq = {"id": "FR1.C1", "module": "FR1", "title_az": "Real ÜDM tənliyi", "estimator": "OLS", "sample": {"n": 20},
              "fit": {"r2": 0.9, "r2_adj": 0.88}, "used_in_forecast": True, "dependent": {"code": "ln_rgdp", "label_az": "ln real ÜDM"},
              "coefficients": [{"name": "c", "coef": 1.0, "se": 0.1, "t": 10, "p": 0.0, "ci_low": 0.8, "ci_high": 1.2}],
              "fitted": {"years": [2000], "actual": [1], "fitted": [1], "resid": [0]}, "summary_text": "...",
              "robustness": {"verdict": "stabil", "recursive": {"years": [2010]}}}
        (out / "FR1_equations.json").write_text(json.dumps({"module": "FR1", "generated": "2026-10-05T00:00:00Z",
                                                             "data_mode": "OBSERVED", "equations": [eq]}), encoding="utf-8")
        cls.s = H.ServerProc(cls.root, "--engine-chain", "stub_engine", "--engine-path", str(H.TESTS))

    @classmethod
    def tearDownClass(cls):
        cls.s.stop()
        assert sorted((p.name, p.stat().st_mtime_ns) for p in cls.real_fp.iterdir()) == cls.real_state, "real data/ changed!"
        shutil.rmtree(cls.root, ignore_errors=True)

    # ------------------------------------------------------------------ basics
    def test_health_no_token(self):
        code, d = self.s.req("GET", "/api/v1/health", token=None)
        self.assertEqual(code, 200)
        self.assertEqual(d["status"], "ok")
        code, _, body = self.s.req("GET", "/api/v1/openapi.yaml", token=None, raw=True)
        self.assertEqual(code, 200)
        self.assertIn(b"openapi: 3.0", body)

    def test_auth_enforced(self):
        self.assertEqual(self.s.req("GET", "/api/v1/status", token=None)[0], 401)
        code, d = self.s.req("GET", "/api/v1/status", token="wrong")
        self.assertEqual(code, 401)
        self.assertIn("Nişan", d["error"]["message"])
        self.assertEqual(self.s.req("GET", "/api/v1/status", token=H.READ)[0], 200)
        self.assertEqual(self.s.req("GET", "/api/v1/status", token=H.WRITE)[0], 200)        # write includes read
        code, d = self.s.req("POST", "/api/v1/runs", {"only": "FR5"}, token=H.READ)
        self.assertEqual(code, 403)
        self.assertIn("icazəsi yoxdur", d["error"]["message"])
        self.assertEqual(self.s.req("DELETE", "/api/v1/scenarios/saved/x", token=H.READ)[0], 403)
        self.assertEqual(self.s.req("POST", "/api/v1/uploads", b"x", token=None, headers={"X-Filename": "a.csv"})[0], 401)

    def test_cors_localhost_only(self):
        code, h, _ = self.s.req("GET", "/api/v1/health", token=None, raw=True, headers={"Origin": "http://localhost:5173"})
        self.assertEqual(h.get("Access-Control-Allow-Origin"), "http://localhost:5173")
        code, h, _ = self.s.req("GET", "/api/v1/health", token=None, raw=True, headers={"Origin": "https://evil.example"})
        self.assertNotIn("Access-Control-Allow-Origin", h)
        code, h, _ = self.s.req("OPTIONS", "/api/v1/uploads", token=None, raw=True, headers={"Origin": "http://127.0.0.1:8790"})
        self.assertEqual(code, 204)
        self.assertIn("X-Filename", h.get("Access-Control-Allow-Headers", ""))

    def test_static_files(self):
        for path, ctype in (("/", "text/html"), ("/panel/", "text/html"), ("/site/", "text/html"),
                            ("/panel/assets/panel.css", "text/css")):
            code, h, body = self.s.req("GET", path, token=None, raw=True)
            self.assertEqual(code, 200, path)
            self.assertTrue(h["Content-Type"].startswith(ctype), (path, h["Content-Type"]))
        code, h, _ = self.s.req("GET", "/panel/data/fr4.js", token=None, raw=True)
        self.assertTrue(h["Content-Type"].startswith("application/javascript"))
        self.assertEqual(self.s.req("GET", "/panel/assets/", token=None)[0], 404)            # no directory listing
        self.assertEqual(self.s.req("GET", "/data/secret.txt", token=None)[0], 404)          # outside the trees

    def test_path_traversal_blocked(self):
        link = self.root / "panel" / "escape"
        try:
            os.symlink(self.root / "data", link)
        except OSError:
            link = None
        evil = ["/panel/../data/secret.txt", "/panel/%2e%2e/data/secret.txt", "/panel/..%2fdata%2fsecret.txt",
                "/site/%2e%2e%2f%2e%2e%2fetc%2fpasswd", "/panel/%2E%2E/%2E%2E/etc/passwd", "/panel/..\\data\\secret.txt",
                "/panel/%5c..%5cdata%5csecret.txt", "/panel/_build/pcore.py", "/panel/build_panel.py", "/panel/.hidden",
                "/panel/%00", "/panel/%252e%252e/data/secret.txt", "/site/../../../../etc/passwd", "/panel//etc/passwd",
                "/panel/escape/secret.txt", "/panel/escape/"]
        for p in evil:
            code, body = H.raw_http(self.s.base, p)
            self.assertIn(code, (400, 403, 404), p)
            self.assertNotIn(b"TOP-SECRET", body, p)
            self.assertNotIn(b"SECRET_SOURCE", body, p)
            self.assertNotIn(b"root:", body, p)

    # ------------------------------------------------------------------ uploads
    def upload_raw(self, name, data, kind=None, token=H.WRITE, query="", **hdr):
        h = {"X-Filename": name, "Content-Type": "application/octet-stream", **hdr}
        if kind:
            h["X-Kind"] = kind
        return self.s.req("POST", "/api/v1/uploads" + query, data, token=token, headers=h)

    def test_firm_panel_broken_rejected_with_row_errors(self):
        fields, rows = H.firm_panel_rows(60)
        rows[3]["cash"] = "-5"
        rows[7]["region"] = "Atlantis"
        rows[9]["firm_id"], rows[9]["year"] = rows[8]["firm_id"], rows[8]["year"]
        rows[11]["total_assets"] = "abc"
        code, d = self.upload_raw("FR10_firm_panel.csv", H.to_csv_bytes(fields, rows), "firm_panel")
        self.assertEqual(code, 422, d)
        v = d["validation"]
        self.assertFalse(v["ok"])
        got = {(e["row"], e["field"]) for e in v["errors"]}
        self.assertTrue({(5, "cash"), (9, "region"), (10, "firm_id"), (11, "firm_id"), (13, "total_assets")} <= got, got)
        self.assertTrue(all(e.get("message") for e in v["errors"]))
        self.assertIn("mənfi dəyər", {e["message"] for e in v["errors"]})
        code, d2 = self.s.req("POST", "/api/v1/uploads/%s/apply" % d["upload"]["id"], token=H.WRITE)
        self.assertEqual(code, 409)
        self.assertFalse((self.root / "data/firm_panel/FR10_firm_panel.csv").exists())

    def test_firm_panel_synthetic_marker_rejected(self):
        fields, rows = H.firm_panel_rows(20, real=False)
        code, d = self.upload_raw("FR10_firm_panel.csv", H.to_csv_bytes(fields, rows), "firm_panel")
        self.assertEqual(code, 422)
        self.assertIn("sintetik", d["validation"]["errors"][0]["message"])

    def test_firm_panel_valid_applied_to_temp_data(self):
        syn = self.root / "data/firm_panel/FR10_firm_panel_SYNTHETIC.csv"
        syn_sha = sha(syn)
        fields, rows = H.firm_panel_rows(120)
        body = H.to_csv_bytes(fields, rows)
        code, d = self.upload_raw("FR10_firm_panel_2026.csv", body, "firm_panel")
        self.assertEqual(code, 201, d)
        self.assertEqual(d["validation"]["summary"]["mode"], "REAL")
        self.assertEqual(d["validation"]["target"], "data/firm_panel/FR10_firm_panel.csv")
        self.assertEqual(d["validation"]["affected_stages"], ["FR10", "FR12"])
        uid = d["upload"]["id"]
        self.assertTrue((self.root / "data/inbox" / uid / "FR10_firm_panel_2026.csv").exists())
        code, a = self.s.req("POST", "/api/v1/uploads/%s/apply" % uid, token=H.WRITE)
        self.assertEqual(code, 200, a)
        tgt = self.root / "data/firm_panel/FR10_firm_panel.csv"
        self.assertEqual(sha(tgt), hashlib.sha256(body).hexdigest())
        self.assertEqual(sha(syn), syn_sha)                                           # synthetic file untouched
        self.assertEqual(self.s.req("POST", "/api/v1/uploads/%s/apply" % uid, token=H.WRITE)[0], 409)  # twice
        code, st = self.s.req("GET", "/api/v1/status")
        self.assertEqual(st["data_modes"]["FR10"]["input_mode"], "REAL")
        self.assertTrue(st["data_modes"]["FR10"]["rerun_needed"])
        # a second panel replaces the first one: the first goes to data/_replaced/
        rows[0]["revenue"] = str(float(rows[0]["revenue"]) + 1)
        code, d = self.upload_raw("FR10_firm_panel.csv", H.to_csv_bytes(fields, rows), "firm_panel", query="?apply=1")
        self.assertEqual(code, 201, d)
        rep = d["apply"]
        self.assertEqual(len(rep["replaced"]), 1)
        self.assertTrue((self.root / rep["replaced"][0]["backup"]).exists())
        self.assertEqual(rep["replaced"][0]["sha256"], hashlib.sha256(body).hexdigest())

    def test_multipart_business_register(self):
        def mut(rows):
            for r in rows:
                r["data_status"] = "Ministry register"
        H.head_csv(H.REAL_ROOT / "data/business_register/FR12_business_register_SYNTHETIC.csv", self.root / "br.csv", 150, mut)
        data = (self.root / "br.csv").read_bytes()
        body, ctype = multipart_form.build({"kind": "business_register", "note": "sınaq"},
                                           {"file": ("FR12_business_register_Ə.csv", data, "text/csv")})
        code, d = self.s.req("POST", "/api/v1/uploads", body, token=H.WRITE, headers={"Content-Type": ctype})
        self.assertEqual(code, 201, d)
        self.assertEqual(d["upload"]["kind"], "business_register")
        self.assertEqual(d["upload"]["filename"], "FR12_business_register_Ə.csv")
        self.assertEqual(d["validation"]["target"], "data/business_register/FR12_business_register.csv")
        self.assertEqual(d["validation"]["affected_stages"], ["FR12"])
        bad, ctype = multipart_form.build({"kind": "business_register"}, {})
        self.assertEqual(self.s.req("POST", "/api/v1/uploads", bad, token=H.WRITE, headers={"Content-Type": ctype})[0], 400)
        code, d = self.s.req("POST", "/api/v1/uploads", b"garbage", token=H.WRITE,
                             headers={"Content-Type": "multipart/form-data; boundary=zzz"})
        self.assertEqual(code, 400)
        self.assertIn("multipart", d["error"]["message"])

    def test_dsk_upload(self):
        code, d = self.upload_raw("002_1-2en.xls", b"<html>not published</html>", "dsk")
        self.assertEqual(code, 422)
        self.assertIn("BIFF", d["validation"]["errors"][0]["message"])
        good = (H.REAL_ROOT / "data/dsk/002_1-2en.xls").read_bytes()
        code, d = self.upload_raw("002_1-2en.xls", good, "dsk")                     # subfolder inferred
        self.assertEqual(code, 201, d)
        self.assertEqual(d["validation"]["subfolder"], "dsk")
        self.assertEqual(d["validation"]["affected_stages"], ["FR4", "FR10", "FR12"])
        code, d = self.upload_raw("002_1-2en.xls", good, "dsk", **{"X-Subfolder": "../../etc"})
        self.assertEqual(code, 422)
        code, d = self.upload_raw("unknown_table.xls", good, "dsk")
        self.assertEqual(code, 422)
        self.assertIn("subfolder", d["validation"]["errors"][0]["message"])

    def test_filename_sanitised_and_kind_inferred(self):
        fields, rows = H.firm_panel_rows(30)
        code, d = self.upload_raw("../../../FR10_firm_panel.csv", H.to_csv_bytes(fields, rows))   # no X-Kind: inferred
        self.assertEqual(code, 201, d)
        self.assertEqual(d["upload"]["kind"], "firm_panel")
        self.assertEqual(d["upload"]["filename"], "FR10_firm_panel.csv")
        stored = (self.root / d["upload"]["stored_path"]).resolve()
        self.assertTrue(str(stored).startswith(str((self.root / "data/inbox").resolve())))
        self.assertEqual(self.upload_raw("notes.docx", b"x")[0], 400)
        self.assertEqual(self.upload_raw("a.csv", b"x", "virus")[0], 400)
        self.assertEqual(self.upload_raw("a.csv", b"", "firm_panel")[0], 400)

    def test_series_upload_and_observations(self):
        body = json.dumps({"actor": "dsk", "source": "DSK 2025", "items": [
            {"id": "fr1:rgdp", "period": 2025, "value": 101.5}, {"module": "FR1", "code": "cpi", "period": 2025, "value": 2.9}]}).encode()
        code, d = self.upload_raw("obs.json", body, "series", query="?apply=1")
        self.assertEqual(code, 201, d)
        self.assertEqual(d["apply"]["inserted"], 2)
        code, o = self.s.req("GET", "/api/v1/observations?module=FR1")
        self.assertEqual(o["total"], 2)
        self.assertEqual({(x["code"], x["value"]) for x in o["items"]}, {("rgdp", 101.5), ("cpi", 2.9)})
        bad = json.dumps({"items": [{"id": "fr1:nope", "period": 2025, "value": 1}, {"id": "fr1:rgdp", "period": 1800, "value": 1},
                                    {"id": "fr1:rgdp", "period": 2025, "value": "x"}]}).encode()
        code, d = self.upload_raw("obs.json", bad, "series")
        self.assertEqual(code, 422)
        self.assertEqual({e["reason"] for e in d["validation"]["errors"]}, {"unknown_series", "period_out_of_range", "bad_item"})

    def test_uploads_list_and_get(self):
        code, d = self.s.req("GET", "/api/v1/uploads?limit=5")
        self.assertEqual(code, 200)
        self.assertIsInstance(d["items"], list)
        self.assertEqual(self.s.req("GET", "/api/v1/uploads/u-none")[0], 404)

    # ------------------------------------------------------------------ data
    def test_forecasts_catalog_and_panel_fallback(self):
        code, d = self.s.req("GET", "/api/v1/forecasts?module=FR1&scenario=Baseline&from=2026")
        self.assertEqual(code, 200, d)
        pts = {(x["id"], x["period"]): x["value"] for x in d["items"]}
        self.assertEqual(pts[("fr1:rgdp", 2027)], 104.0)
        self.assertEqual(d["series"]["fr1:rgdp"]["label_az"], "Real ÜDM")
        code, d = self.s.req("GET", "/api/v1/forecasts?id=fr1:cpi&scenario=Adverse")
        self.assertEqual([x["value"] for x in d["items"]], [3.5, 3.6])
        code, d = self.s.req("GET", "/api/v1/forecasts?id=fr4:emp:agr&scenario=Baseline,Actual")
        self.assertEqual(d["source"], "panel")
        self.assertEqual([x["value"] for x in d["items"] if x["scenario"] == "Baseline"], [1, 2, 3, 4, 5])
        self.assertEqual(self.s.req("GET", "/api/v1/forecasts")[0], 400)
        self.assertEqual(self.s.req("GET", "/api/v1/forecasts?module=FR99")[0], 400)
        self.assertEqual(self.s.req("GET", "/api/v1/forecasts?id=fr1:zzz")[0], 404)
        code, c = self.s.req("GET", "/api/v1/catalog?module=FR1")
        self.assertEqual(c["sources"]["FR1"], "catalog")
        self.assertEqual(c["total"], 2)

    def test_equations(self):
        code, d = self.s.req("GET", "/api/v1/equations?module=FR1")
        self.assertEqual(code, 200)
        self.assertEqual(d["items"][0]["id"], "FR1.C1")
        self.assertNotIn("fitted", d["items"][0])                                  # compact list
        self.assertEqual(d["items"][0]["robustness"], {"verdict": "stabil"})
        code, d = self.s.req("GET", "/api/v1/equations?id=FR1.C1")
        self.assertIn("fitted", d["items"][0])
        code, d = self.s.req("GET", "/api/v1/equations?module=FR3")
        self.assertEqual((code, d["missing"]), (200, ["FR3"]))
        self.assertIn("hələ yaradılmayıb", d["message"])
        self.assertEqual(self.s.req("GET", "/api/v1/equations?id=FR1.nope")[0], 404)

    def test_report_xlsx(self):
        import openpyxl
        code, h, body = self.s.req("POST", "/api/v1/reports", {"title": "Sınaq", "indicators": ["fr1:rgdp", "fr1:cpi"],
                                                               "scenarios": ["Baseline", "Adverse"], "equations": True},
                                   token=H.WRITE, raw=True)
        self.assertEqual(code, 200, body[:300])
        self.assertTrue(body.startswith(b"PK"))
        self.assertIn("attachment", h["Content-Disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(body))
        self.assertEqual(wb.sheetnames, ["Məlumat", "Göstəricilər", "Tənliklər"])
        ws = wb["Göstəricilər"]
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(rows[0][:5], ("ID", "Modul", "Göstərici", "Vahid", "Ssenari"))
        base = [r for r in rows if r[0] == "fr1:rgdp" and r[4] == "Əsas"][0]
        self.assertEqual(base[rows[0].index(2027)], 104.0)
        self.assertEqual(wb["Tənliklər"]["A2"].value, "FR1.C1")
        self.assertEqual(self.s.req("POST", "/api/v1/reports", {"scenarios": ["X"]}, token=H.WRITE)[0], 400)

    # ------------------------------------------------------------------ scenarios (stub engine)
    def test_scenarios_with_stub_engine(self):
        code, d = self.s.req("GET", "/api/v1/scenarios/inputs?module=fr1")
        self.assertEqual(code, 200, d)
        self.assertEqual(d["inputs"]["exogenous"][0]["id"], "fr1:oil_price")
        self.assertEqual(self.s.req("GET", "/api/v1/scenarios/inputs?module=FR99")[0], 400)
        code, base = self.s.req("POST", "/api/v1/scenarios/run", {"overrides": {}, "scenario": "Baseline"}, token=H.WRITE)
        self.assertEqual(code, 200, base)
        code, shock = self.s.req("POST", "/api/v1/scenarios/run", {
            "overrides": {"FR1": {"exogenous": {"fr1:oil_price": {"pct": 20}}}}, "scenario": "Baseline", "name": "Neft +20%",
            "author": "analitik"}, token=H.WRITE)
        self.assertEqual(code, 200, shock)
        g0 = base["result"]["results"]["FR1"]["series"]["fr1:rgdp"]["2030"]
        g1 = shock["result"]["results"]["FR1"]["series"]["fr1:rgdp"]["2030"]
        self.assertGreater(g1, g0)
        self.assertIn("FR12", shock["result"]["results"])                          # chained downstream
        sid = shock["saved"]["id"]
        code, lst = self.s.req("GET", "/api/v1/scenarios/saved")
        self.assertIn(sid, [x["id"] for x in lst["items"]])
        code, one = self.s.req("GET", "/api/v1/scenarios/saved/%s" % sid)
        self.assertEqual((one["name"], one["author"]), ("Neft +20%", "analitik"))
        self.assertEqual(one["overrides"]["FR1"]["exogenous"]["fr1:oil_price"]["pct"], 20)
        code, s2 = self.s.req("POST", "/api/v1/scenarios/saved", {"name": "Əl ilə", "overrides": {"FR4": {"levers": {}}}}, token=H.WRITE)
        self.assertEqual(code, 201)
        code, s3 = self.s.req("POST", "/api/v1/scenarios/saved", {"id": s2["id"], "name": "Yeniləndi"}, token=H.WRITE)
        self.assertEqual((code, s3["name"], s3["created_at"]), (200, "Yeniləndi", s2["created_at"]))
        self.assertEqual(self.s.req("DELETE", "/api/v1/scenarios/saved/%s" % sid, token=H.WRITE)[0], 200)
        self.assertEqual(self.s.req("DELETE", "/api/v1/scenarios/saved/%s" % sid, token=H.WRITE)[0], 404)
        code, d = self.s.req("POST", "/api/v1/scenarios/run", {"overrides": {"FR9": {}}}, token=H.WRITE)
        self.assertEqual(code, 400)
        code, d = self.s.req("POST", "/api/v1/scenarios/run", {"overrides": {"FR1": {"bogus": 1}}}, token=H.WRITE)
        self.assertEqual(code, 400)
        code, d = self.s.req("POST", "/api/v1/scenarios/run", {"overrides": {"FR1": {"exogenous": {"fr1:oil_price": [1, 2]}}}},
                             token=H.WRITE)
        self.assertEqual(code, 400)                                                  # engine ValueError -> 400
        self.assertIn("qəbul edilmədi", d["error"]["message"])
        code, d = self.s.req("POST", "/api/v1/scenarios/run", {"scenario": "Nope"}, token=H.WRITE)
        self.assertEqual(code, 400)
        code, d = self.s.req("POST", "/api/v1/scenarios/run", {"modules": ["FR1"], "scenario": "Adverse"}, token=H.WRITE)
        self.assertEqual(list(d["result"]["results"]), ["FR1"])

    def test_status(self):
        code, d = self.s.req("GET", "/api/v1/status")
        self.assertEqual(code, 200)
        self.assertEqual(d["data_modes"]["FR10"]["mode"], "SYNTHETIC")             # outputs are FR10_SYNTHETIC_*
        self.assertEqual(d["data_modes"]["FR12"]["input_mode"], "SYNTHETIC")       # no real register applied
        self.assertIsNone(d["data_modes"]["FR12"]["output_mode"])                  # no FR12 outputs in the temp root
        self.assertTrue(d["engine"]["available"])
        self.assertTrue(d["catalogs"]["FR1"])
        self.assertFalse(d["catalogs"]["FR3"])
        self.assertIn("workbook", d["data_vintage"])
        self.assertEqual(self.s.req("GET", "/api/v1/nonexistent")[0], 404)


class WorkbookUploadTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = H.make_root(with_workbook=True)
        cls.wb = next((cls.root / "data").glob("Statistik data dinamika*.xlsx"))
        cls.s = H.ServerProc(cls.root)

    @classmethod
    def tearDownClass(cls):
        cls.s.stop()
        shutil.rmtree(cls.root, ignore_errors=True)

    def test_workbook_broken_then_valid_applied(self):
        import openpyxl
        from urllib.parse import quote
        wb = openpyxl.load_workbook(self.wb)
        wb["Real sektor"]["A29"] = "Yanlış sətir"
        buf = io.BytesIO()
        wb.save(buf)
        name = "Statistik data dinamika 05.09.2026.xlsx"
        h = {"X-Filename": quote(name), "X-Kind": "workbook", "Content-Type": "application/octet-stream"}
        code, d = self.s.req("POST", "/api/v1/uploads", buf.getvalue(), token=H.WRITE, headers=h)
        self.assertEqual(code, 422, d)
        e = [x for x in d["validation"]["errors"] if x.get("module") == "FR1"][0]
        self.assertEqual((e["sheet"], e["row"], e["found"]), ("Real sektor", 29, "Yanlış sətir"))
        old = sha(self.wb)
        good = self.wb.read_bytes()
        code, d = self.s.req("POST", "/api/v1/uploads?apply=1", good, token=H.WRITE, headers=h)
        self.assertEqual(code, 201, d.get("validation", {}).get("errors"))
        self.assertEqual(d["upload"]["filename"], name)
        self.assertEqual(d["validation"]["target"], "data/" + self.wb.name)               # the exact name the notebooks read
        self.assertEqual(d["validation"]["affected_stages"], ["FR1", "FR3", "FR4", "FR5", "FR10", "FR12"])
        self.assertGreater(d["validation"]["summary"]["addresses_checked"], 250)
        self.assertEqual(sha(self.wb), old)
        self.assertTrue((self.root / d["apply"]["backup"] / self.wb.name).exists())
        code, st = self.s.req("GET", "/api/v1/status")
        self.assertEqual(st["data_vintage"]["workbook"]["original_filename"], name)
        self.assertEqual(st["data_vintage"]["workbook"]["last_annual_year"], 2025)
        self.assertEqual(self.s.req("GET", "/panel", token=None)[0], 200)                # redirect to /panel/ followed


class EngineUnavailableTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = H.make_root()
        cls.s = H.ServerProc(cls.root, "--engine-chain", "nonexistent_engine_xyz.chain")

    @classmethod
    def tearDownClass(cls):
        cls.s.stop()
        shutil.rmtree(cls.root, ignore_errors=True)

    def test_503_in_azerbaijani(self):
        code, d = self.s.req("GET", "/api/v1/scenarios/inputs?module=FR1")
        self.assertEqual(code, 503)
        self.assertEqual(d["error"]["code"], "engine_unavailable")
        self.assertIn("mühərriki hələ mövcud deyil", d["error"]["message"])
        self.assertEqual(self.s.req("POST", "/api/v1/scenarios/run", {"overrides": {}}, token=H.WRITE)[0], 503)
        code, st = self.s.req("GET", "/api/v1/status")
        self.assertFalse(st["engine"]["available"])
        code, sv = self.s.req("POST", "/api/v1/scenarios/saved", {"name": "mühərriksiz"}, token=H.WRITE)
        self.assertEqual(code, 201)                                                # saved scenarios still work


if __name__ == "__main__":
    unittest.main()
